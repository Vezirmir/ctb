from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
WRAP = Alignment(wrap_text=True, vertical="top")


HEADER_FILL = PatternFill("solid", fgColor="E6E5E0")
REQUIRED_FILL = PatternFill("solid", fgColor="FFC9A8")


def _company_values(c):
    contacts = "\n".join(
        " / ".join(part for part in (p.full_name, p.position, p.email, p.phone) if part)
        for p in c.contacts.all()
    )
    return {
        "name": c.name,
        "country": str(c.country) if c.country else "",
        "industries": ", ".join(str(i) for i in c.industries.all()),
        "city": c.city,
        "websites": c.websites,
        "phones": c.phones,
        "emails": "\n".join(e.email for e in c.emails.all()),
        "contacts": contacts,
        "address": c.address,
        "tags": ", ".join(t.name for t in c.tags.all()),
        "description": c.description,
        "notes": c.notes,
    }


def _write_header(ws, with_hints=False):
    from .importer import TEMPLATE_COLUMNS

    ws.append([str(header) for _field, header, _width, _hint in TEMPLATE_COLUMNS])
    for index, (field_name, _header, width, hint) in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = REQUIRED_FILL if field_name == "name" else HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if with_hints and hint:
            cell.comment = Comment(str(hint), "CTB")
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def companies_to_xlsx_response(queryset):
    from .importer import TEMPLATE_COLUMNS

    queryset = queryset.select_related("country").prefetch_related(
        "industries", "tags", "contacts", "emails")
    wb = Workbook()
    ws = wb.active
    ws.title = _("Companies")[:31]
    _write_header(ws)
    for c in queryset:
        values = _company_values(c)
        ws.append([values[field_name] for field_name, *_rest in TEMPLATE_COLUMNS])
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP
    ws.auto_filter.ref = ws.dimensions
    return _xlsx_response(wb, f"companies_{timezone.localdate():%Y-%m-%d}.xlsx")


def import_template_response():
    """Empty company list to fill in and import, with drop-down lists and instructions."""
    from .importer import TEMPLATE_COLUMNS
    from .models import Country, Industry

    wb = Workbook()
    ws = wb.active
    ws.title = _("Companies")[:31]
    _write_header(ws, with_hints=True)
    for row in range(2, 502):
        ws.row_dimensions[row].height = 30
        for col in range(1, len(TEMPLATE_COLUMNS) + 1):
            ws.cell(row=row, column=col).alignment = WRAP

    # Drop-down lists (typing other values is still allowed).
    lists = wb.create_sheet("lists")
    lists.sheet_state = "hidden"
    columns = {field_name: get_column_letter(i)
               for i, (field_name, *_rest) in enumerate(TEMPLATE_COLUMNS, start=1)}
    for list_col, field_name, values in (
        ("A", "country", sorted(str(c) for c in Country.objects.all())),
        ("B", "industries", sorted(str(i) for i in Industry.objects.all())),
    ):
        for row, value in enumerate(values, start=1):
            lists[f"{list_col}{row}"] = value
        if values:
            validation = DataValidation(
                type="list", formula1=f"=lists!${list_col}$1:${list_col}${len(values)}",
                allow_blank=True, showErrorMessage=False)
            ws.add_data_validation(validation)
            validation.add(f"{columns[field_name]}2:{columns[field_name]}501")

    help_sheet = wb.create_sheet(_("How to fill in")[:31])
    help_sheet.column_dimensions["A"].width = 110
    lines = [
        (_("How to fill in the template"), True),
        (_("1. Fill in one company per row on the first sheet. Only “Company name” is required."), False),
        (_("2. Several websites, phones or e-mails: one per line (Alt+Enter) or separated by commas."), False),
        (_("3. Contact persons: one per line, as “Full name / Position / e-mail / phone”."), False),
        (_("4. Country and industries can be picked from the list or typed in."), False),
        (_("5. Save the file and upload it on the “Import from Excel” page. Existing companies are updated, not duplicated."), False),
    ]
    for text, bold in lines:
        help_sheet.append([text])
        help_sheet.cell(row=help_sheet.max_row, column=1).font = Font(bold=bold, size=13 if bold else 11)
    help_sheet.append([])
    help_sheet.append([_("Example:")])
    help_sheet.cell(row=help_sheet.max_row, column=1).font = Font(bold=True)
    example = {
        "name": "Example Trade LLC", "country": "Türkiye", "industries": "Automotive",
        "city": "Bursa", "websites": "www.example.com", "phones": "+90 224 000 00 00",
        "emails": "info@example.com\nsales@example.com",
        "contacts": "Ali Yılmaz / Purchasing manager / ali@example.com / +90 532 000 00 00",
        "address": "", "tags": "VIP", "description": "", "notes": "",
    }
    for field_name, header, _width, _hint in TEMPLATE_COLUMNS:
        help_sheet.append([f"{header}: {example[field_name]}".replace("\n", "; ")])

    return _xlsx_response(wb, "ctb_import_template.xlsx")


def _xlsx_response(wb, filename):
    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def _local(dt):
    return timezone.localtime(dt) if dt else None


def schedule_to_xlsx_response(event):
    from .matchmaking import company_agenda, schedule_rows

    wb = Workbook()
    ws = wb.active
    ws.title = _("Schedule")[:31]
    columns = [(_("Date"), 12), (_("Time"), 8), (_("Table"), 7), (_("First company"), 34),
               (_("Country"), 14), (_("Second company"), 34), (_("Country"), 14), (_("Status"), 14)]
    ws.append([title for title, _width in columns])
    for index, (_title, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for m in schedule_rows(event):
        at = _local(m.scheduled_at)
        ws.append([
            at.strftime("%d.%m.%Y") if at else "", at.strftime("%H:%M") if at else "",
            m.table or "", m.company_a.name, str(m.company_a.country or ""),
            m.company_b.name, str(m.company_b.country or ""), m.get_status_display(),
        ])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    agenda = wb.create_sheet(_("By company")[:31])
    for index, width in enumerate((12, 8, 7, 36, 16, 40), start=1):
        agenda.column_dimensions[get_column_letter(index)].width = width
    for company, items in company_agenda(event).items():
        agenda.append([company.name])
        agenda[agenda.max_row][0].font = Font(bold=True, size=13)
        agenda.append([_("Date"), _("Time"), _("Table"), _("Partner"), _("Country"), _("E-mails")])
        for cell in agenda[agenda.max_row]:
            cell.font = Font(bold=True)
        for m, partner in items:
            at = _local(m.scheduled_at)
            agenda.append([
                at.strftime("%d.%m.%Y") if at else "", at.strftime("%H:%M") if at else "",
                m.table or "", partner.name, str(partner.country or ""),
                ", ".join(partner.emails.values_list("email", flat=True)[:3]),
            ])
        agenda.append([])

    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="schedule_{event.pk}.xlsx"'
    wb.save(response)
    return response
