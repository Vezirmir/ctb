"""Import companies from the CTB Excel lists.

Expected layout (as in the shared archive)::

    <any root>/<Industry>/<Country>/<file>.xlsx   company lists
    <any root>/<file>.xlsx                         e-mails not linked to a company

Columns are recognised by their header, so their order does not matter.
Re-importing the same files is safe: companies are matched by name and country,
e-mails and contact persons are merged instead of duplicated.
"""

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path, PurePosixPath

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import translation
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy as _lazy
from openpyxl import load_workbook

from . import reference
from .models import Company, Contact, Country, Email, Industry, Tag, normalize_name

EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
MAX_MEMBER_SIZE = 50 * 1024 * 1024
MAX_ARCHIVE_SIZE = 200 * 1024 * 1024
HEADER_SCAN_ROWS = 5

# Header prefixes (lower case) for each field. The longest matching prefix wins,
# so "e-mail с сайта ..." is not mistaken for the plain "e-mail" column.
COMPANY_COLUMNS = {
    "name": ["название компании", "компания", "название", "company name", "company", "name",
             "firma adı", "firma"],
    "country": ["страна", "country", "ülke"],
    "industries": ["отрасль", "отрасли", "industry", "industries", "sektör", "sektörler"],
    "websites": ["сайт", "website", "websites", "web sitesi"],
    "phones": ["номера телефонов", "телефон", "phone", "phones", "telefon"],
    "emails": ["почтовые адреса (e-mail)", "почта", "e-mail", "email", "e-mails", "e-posta"],
    "questionnaire": ["сверка с анкетой"],
    "website_check": ["проверка сайта"],
    "website_emails": ["e-mail с сайта"],
    "people": ["руководство и закупки"],
    "ctb_emails": ["e-mail из почты ctb"],
    "correspondence": ["переписка в почте ctb"],
}
UNASSIGNED_COLUMNS = {
    "group": ["группа", "group", "grup"],
    "email": ["e-mail", "email", "e-posta"],
    "person": ["имя"],
    "correspondence": ["что есть в почте"],
    "notes": ["комментарий", "comment"],
}

# Columns of the downloadable import template (and of the Excel export), in order:
# (field, header, column width, hint shown as a cell comment).
TEMPLATE_COLUMNS = [
    ("name", _lazy("Company name"), 34, _lazy("Required.")),
    ("country", _lazy("Country"), 16, _lazy("Pick from the list or type a new country.")),
    ("industries", _lazy("Industries"), 22,
     _lazy("Pick from the list; several industries can be separated by commas.")),
    ("city", _lazy("City"), 16, ""),
    ("websites", _lazy("Websites"), 26, _lazy("One per line or separated by commas.")),
    ("phones", _lazy("Phones"), 22, _lazy("One per line or separated by commas.")),
    ("emails", _lazy("E-mails"), 30, _lazy("One per line or separated by commas.")),
    ("contacts", _lazy("Contact persons"), 44,
     _lazy("One person per line: Full name / Position / e-mail / phone.")),
    ("address", _lazy("Address"), 30, ""),
    ("tags", _lazy("Tags"), 18, _lazy("Separated by commas.")),
    ("description", _lazy("Description"), 36, ""),
    ("notes", _lazy("Notes"), 36, ""),
]


def company_column_spec():
    """COMPANY_COLUMNS plus the template headers in every interface language."""
    spec = {name: list(aliases) for name, aliases in COMPANY_COLUMNS.items()}
    for code, _name in settings.LANGUAGES:
        with translation.override(code):
            for field_name, header, _width, _hint in TEMPLATE_COLUMNS:
                alias = " ".join(str(header).split()).lower()
                if alias not in spec.setdefault(field_name, []):
                    spec[field_name].append(alias)
    return spec


EMAIL_RE = re.compile(r"[\w.+'-]+@[\w-]+(?:\.[\w-]+)+")
EMAIL_LINE_RE = re.compile(
    r"^\s*(?P<email>[^\s()]+@[^\s()]+)\s*(?:\((?P<person>[^)]*)\))?\s*(?:[—–]\s*(?P<rest>.*))?$"
)
PERSON_LINE_RE = re.compile(r"^(?P<category>[^:]{1,60}):\s*(?P<body>.*?)\s+[—–]\s+(?P<tail>.*)$")
SENT_RE = re.compile(r"писем от нас:\s*(\d+)(?:,\s*посл\.\s*(\d{2}\.\d{2}\.\d{4}))?")


@dataclass
class ImportReport:
    files: list = field(default_factory=list)
    companies_created: int = 0
    companies_updated: int = 0
    emails_created: int = 0
    contacts_created: int = 0
    unassigned_emails_created: int = 0
    warnings: list = field(default_factory=list)

    def summary(self):
        return _(
            "Files: %(files)d. Companies: %(created)d new, %(updated)d updated. "
            "E-mails: %(emails)d new. Contact persons: %(contacts)d new. "
            "E-mails without company: %(unassigned)d new."
        ) % {
            "files": len(self.files),
            "created": self.companies_created,
            "updated": self.companies_updated,
            "emails": self.emails_created,
            "contacts": self.contacts_created,
            "unassigned": self.unassigned_emails_created,
        }


# --------------------------------------------------------------------------- parsing helpers

def _lines(value):
    if value is None:
        return []
    return [line.strip() for line in str(value).splitlines() if line.strip()]


def _text(value):
    return "\n".join(_lines(value))


def _valid_email(value):
    value = value.strip().strip(".,;").lower()
    try:
        validate_email(value)
    except ValidationError:
        return None
    return value


def parse_email_line(line):
    """'a@b.ru (Name) — rest' -> (email, person, rest)."""
    match = EMAIL_LINE_RE.match(line)
    if match:
        return (
            _valid_email(match["email"]),
            (match["person"] or "").strip(),
            (match["rest"] or "").strip(),
        )
    found = EMAIL_RE.search(line)
    if not found:
        return None, "", line
    return _valid_email(found.group()), "", line.replace(found.group(), "").strip(" —–-")


def parse_correspondence(text):
    """Extract counters from 'писем от нас: 2, посл. 02.09.2022; от адреса: отвечал ...'."""
    sent_count = last_sent_on = None
    match = SENT_RE.search(text)
    if match:
        sent_count = int(match[1])
        if match[2]:
            last_sent_on = datetime.strptime(match[2], "%d.%m.%Y").date()
    lowered = text.lower()
    return {
        "sent_count": sent_count,
        "last_sent_on": last_sent_on,
        "replied": "отвечал" in lowered,
        "needs_review": "проверьте" in lowered or "предположительно" in lowered,
    }


def parse_person_line(line):
    """'Директор: Иванов И.И., Генеральный директор — a@b.ru' -> dict for Contact."""
    match = PERSON_LINE_RE.match(line)
    if match:
        category_text, body, tail = match["category"], match["body"], match["tail"]
    else:
        category_text, body, tail = "", line, ""
    found = EMAIL_RE.search(tail)
    email = _valid_email(found.group()) if found else None

    if ", " in body:
        full_name, position = body.split(", ", 1)
    else:
        full_name, position = "", body

    lowered = category_text.lower()
    if "директор" in lowered or "руковод" in lowered:
        category = Contact.Category.MANAGEMENT
    elif "закуп" in lowered or "снабж" in lowered:
        category = Contact.Category.PROCUREMENT
    elif "продаж" in lowered or "сбыт" in lowered:
        category = Contact.Category.SALES
    else:
        category = Contact.Category.OTHER
    return {
        "full_name": full_name.strip()[:200],
        "position": position.strip()[:300],
        "category": category,
        "email": email or "",
    }


def _map_columns(header, spec):
    aliases = sorted(
        ((alias, name) for name, names in spec.items() for alias in names),
        key=lambda item: len(item[0]),
        reverse=True,
    )
    mapping = {}
    for index, cell in enumerate(header):
        title = " ".join(str(cell or "").split()).lower()
        if not title:
            continue
        for alias, name in aliases:
            matches = (
                title == alias
                or title.startswith((alias + " ", alias + "("))
                or (len(alias) > 6 and title.startswith(alias))
            )
            if matches and name not in mapping:
                mapping[name] = index
                break
    return mapping


def _items(value, separators=r"[\n,;]"):
    """Values separated by new lines, commas or semicolons."""
    if value is None:
        return []
    return [item.strip() for item in re.split(separators, str(value)) if item.strip()]


def parse_contact_line(line):
    """'Full name / Position / e-mail / phone' -> dict for Contact (order of e-mail/phone free)."""
    separator = r"\s+/\s+" if re.search(r"\s/\s", line) else r"\s*[,;]\s*"
    parts = [p.strip() for p in re.split(separator, line) if p.strip()]
    email = phone = ""
    text = []
    for part in parts:
        if not email and EMAIL_RE.fullmatch(part):
            email = _valid_email(part) or ""
        elif not phone and sum(ch.isdigit() for ch in part) >= 6 and not re.search(r"[^\W\d_]", part):
            phone = part
        else:
            text.append(part)
    return {
        "full_name": (text[0] if text else "")[:200],
        "position": " / ".join(text[1:])[:300],
        "category": Contact.Category.OTHER,
        "email": email,
        "phone": phone[:60],
    }


def _merge_lines(existing, new_lines):
    lines = _lines(existing)
    seen = {line.casefold() for line in lines}
    for line in new_lines:
        if line.casefold() not in seen:
            lines.append(line)
            seen.add(line.casefold())
    return "\n".join(lines)


# --------------------------------------------------------------------------- importer

class Importer:
    def __init__(self):
        self.report = ImportReport()
        self._countries = None
        self._industries = None

    # reference data ------------------------------------------------------------------

    def _lookup(self, cache, name):
        key = normalize_name(name)
        return cache.get(key)

    def _index(self, obj, *extra):
        names = [obj.name_en, obj.name_tr, obj.name_ru, *extra]
        return {normalize_name(n): obj for n in names if n}

    def country(self, name):
        name = (name or "").strip()
        if not name:
            return None
        if self._countries is None:
            self._countries = {}
            for country in Country.objects.all():
                self._countries.update(self._index(country, country.iso_code))
        found = self._lookup(self._countries, name)
        if found:
            return found

        key = normalize_name(name)
        for iso, en, tr, ru_names in reference.COUNTRIES:
            if key in {normalize_name(n) for n in (iso, en, tr, *ru_names)}:
                country = Country.objects.filter(iso_code=iso).first() or Country.objects.filter(
                    name_en=en).first()
                if country is None:
                    country = Country.objects.create(iso_code=iso, name_en=en, name_tr=tr,
                                                     name_ru=ru_names[0])
                break
        else:
            country = Country.objects.filter(name_en=name).first() or Country.objects.create(
                name_en=name, name_ru=name)
            self.report.warnings.append(
                _("New country “%(name)s” was created; please add its English and Turkish names.")
                % {"name": name}
            )
        self._countries.update(self._index(country, name, country.iso_code))
        return country

    def industry(self, name):
        name = (name or "").strip()
        if not name:
            return None
        if self._industries is None:
            self._industries = {}
            for industry in Industry.objects.all():
                self._industries.update(self._index(industry))
        found = self._lookup(self._industries, name)
        if found:
            return found

        key = normalize_name(name)
        for en, tr, ru_names in reference.INDUSTRIES:
            if key in {normalize_name(n) for n in (en, tr, *ru_names)}:
                industry, _created = Industry.objects.get_or_create(
                    name_en=en, defaults={"name_tr": tr, "name_ru": ru_names[0]})
                break
        else:
            industry = Industry.objects.filter(name_en=name).first() or Industry.objects.create(
                name_en=name, name_ru=name)
            self.report.warnings.append(
                _("New industry “%(name)s” was created; please add its English and Turkish names.")
                % {"name": name}
            )
        self._industries.update(self._index(industry, name))
        return industry

    # files ---------------------------------------------------------------------------

    def import_named_files(self, files, industry=None, country=None):
        """files: iterable of (relative path, bytes).

        Industry and country come from the folder or file names unless given explicitly.
        """
        with transaction.atomic():
            for path, content in files:
                industry_name, country_name = context_from_path(path)
                self.import_workbook(content, path, industry_name, country_name,
                                     industry=industry, country=country)
        return self.report

    def import_workbook(self, content, label, industry_name=None, country_name=None,
                        industry=None, country=None):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:  # openpyxl raises many different errors for broken files
            self.report.warnings.append(_("%(file)s: not a readable Excel file.") % {"file": label})
            return
        self.report.files.append(label)
        recognised = False
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            for header_index, header in enumerate(rows[:HEADER_SCAN_ROWS]):
                company_columns = _map_columns(header, company_column_spec())
                if "name" in company_columns and len(company_columns) > 1:
                    if industry is None:
                        industry = self.industry(industry_name)
                    if country is None:
                        country = self.country(country_name)
                    self._import_companies(rows[header_index + 1:], company_columns, label,
                                           sheet.title, header_index + 2, industry, country)
                    recognised = True
                    break
                unassigned_columns = _map_columns(header, UNASSIGNED_COLUMNS)
                if {"group", "email"} <= unassigned_columns.keys():
                    self._import_unassigned(rows[header_index + 1:], unassigned_columns)
                    recognised = True
                    break
        workbook.close()
        if not recognised:
            self.report.warnings.append(
                _("%(file)s: no company list found (a “Company name” column is required).")
                % {"file": label}
            )

    # rows ----------------------------------------------------------------------------

    def _import_companies(self, rows, columns, label, sheet, first_row, industry, country):
        def cell(row, name):
            index = columns.get(name)
            return row[index] if index is not None and index < len(row) else None

        for offset, row in enumerate(rows):
            name = " ".join(str(cell(row, "name") or "").split())
            if not name:
                if any(v not in (None, "") for v in row):
                    self.report.warnings.append(
                        _("%(file)s, sheet %(sheet)s, row %(row)d: skipped, company name is empty.")
                        % {"file": label, "sheet": sheet, "row": first_row + offset}
                    )
                continue

            row_country = self.country(_text(cell(row, "country"))) or country
            company = Company.objects.filter(
                normalized_name=normalize_name(name), country=row_country).first()
            if company is None:
                company = Company(name=name, country=row_country)
                self.report.companies_created += 1
            else:
                self.report.companies_updated += 1

            company.websites = _merge_lines(
                company.websites, _items(cell(row, "websites"), r"[\n,;\s]+"))
            company.phones = _merge_lines(company.phones, _items(cell(row, "phones")))
            for field_name, column in (("questionnaire_note", "questionnaire"),
                                       ("website_check_note", "website_check"),
                                       ("city", "city"), ("address", "address"),
                                       ("description", "description"), ("notes", "notes")):
                value = _text(cell(row, column))
                if value:
                    setattr(company, field_name, value)
            company.save()

            industries = [industry] if industry else []
            for industry_name in re.split(r"[,;\n]", str(cell(row, "industries") or "")):
                if industry_name.strip():
                    industries.append(self.industry(industry_name))
            if industries:
                company.industries.add(*industries)
            tags = [Tag.objects.get_or_create(name=name[:100])[0]
                    for name in _items(cell(row, "tags"))]
            if tags:
                company.tags.add(*tags)
            for line in _lines(cell(row, "contacts")):
                self._upsert_contact(company, parse_contact_line(line))

            for line in _lines(cell(row, "emails")):
                for found in EMAIL_RE.findall(line):
                    self._upsert_email(company, _valid_email(found), Email.Source.LIST)
            for line in _lines(cell(row, "website_emails")):
                email, person, rest = parse_email_line(line)
                self._upsert_email(company, email, Email.Source.WEBSITE,
                                   person_name=person, description=rest)
            for column, source in (("ctb_emails", Email.Source.CTB_MAIL),
                                   ("correspondence", Email.Source.LIST)):
                for line in _lines(cell(row, column)):
                    email, person, rest = parse_email_line(line)
                    self._upsert_email(company, email, source, person_name=person,
                                       correspondence=rest)
            for line in _lines(cell(row, "people")):
                self._upsert_contact(company, parse_person_line(line))

    def _upsert_email(self, company, email, source, person_name="", description="",
                      correspondence="", overwrite=True):
        if not email:
            return
        obj, created = Email.objects.get_or_create(
            company=company, email=email, defaults={"source": source})
        if created:
            self.report.emails_created += 1
        if person_name and (overwrite or not obj.person_name):
            obj.person_name = person_name[:200]
        if description and (overwrite or not obj.description):
            obj.description = description[:500]
        if correspondence:
            obj.correspondence = correspondence
            parsed = parse_correspondence(correspondence)
            obj.sent_count = parsed["sent_count"]
            obj.last_sent_on = parsed["last_sent_on"]
            obj.replied = obj.replied or parsed["replied"]
            obj.needs_review = parsed["needs_review"]
        obj.save()

    def _upsert_contact(self, company, data):
        if not (data["full_name"] or data["position"]):
            return
        contact, created = Contact.objects.get_or_create(
            company=company, full_name=data["full_name"], position=data["position"],
            defaults={"category": data["category"], "email": data["email"],
                      "phone": data.get("phone", "")},
        )
        if created:
            self.report.contacts_created += 1
        else:
            changed = False
            for field_name in ("email", "phone"):
                if data.get(field_name) and not getattr(contact, field_name):
                    setattr(contact, field_name, data[field_name])
                    changed = True
            if changed:
                contact.save()
        if data["email"]:
            # Several people may share one mailbox: do not overwrite what is already known.
            self._upsert_email(company, data["email"], Email.Source.WEBSITE,
                               person_name=data["full_name"], description=data["position"],
                               overwrite=False)

    def _import_unassigned(self, rows, columns):
        def cell(row, name):
            index = columns.get(name)
            return row[index] if index is not None and index < len(row) else None

        for row in rows:
            found = EMAIL_RE.search(str(cell(row, "email") or ""))
            email = _valid_email(found.group()) if found else None
            if not email:
                continue
            obj = Email.objects.filter(company__isnull=True, email=email).first()
            if obj is None:
                obj = Email(email=email, source=Email.Source.CTB_MAIL)
                self.report.unassigned_emails_created += 1
            obj.group = _text(cell(row, "group"))[:255]
            obj.person_name = _text(cell(row, "person"))[:200] or obj.person_name
            obj.notes = _text(cell(row, "notes")) or obj.notes
            correspondence = _text(cell(row, "correspondence"))
            if correspondence:
                obj.correspondence = correspondence
                parsed = parse_correspondence(correspondence)
                obj.sent_count = parsed["sent_count"]
                obj.last_sent_on = parsed["last_sent_on"]
                obj.replied = obj.replied or parsed["replied"]
            obj.save()


# --------------------------------------------------------------------------- entry points

def context_from_path(path):
    """Industry and country from '<Industry>/<Country>/file.xlsx' or 'Industry - Country.xlsx'."""
    parts = PurePosixPath(str(path).replace("\\", "/")).parts
    if len(parts) >= 3:
        return parts[-3], parts[-2]
    stem = PurePosixPath(parts[-1]).stem
    if " - " in stem:
        industry_name, country_name = stem.rsplit(" - ", 1)
        return industry_name.strip(), country_name.strip()
    return None, None


def _zip_member_name(info):
    if info.flag_bits & 0x800:
        return info.filename
    raw = info.filename.encode("cp437")
    for encoding in ("utf-8", "cp866"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return info.filename


def _is_excel(name):
    path = PurePosixPath(name)
    return (
        path.suffix.lower() in EXCEL_SUFFIXES
        and not path.name.startswith(("~$", "."))
        and "__MACOSX" not in path.parts
    )


def read_zip(fileobj):
    with zipfile.ZipFile(fileobj) as archive:
        members = [info for info in archive.infolist()
                   if not info.is_dir() and _is_excel(_zip_member_name(info))]
        if sum(info.file_size for info in members) > MAX_ARCHIVE_SIZE:
            raise ValueError(_("The archive is too large."))
        files = []
        for info in members:
            if info.file_size > MAX_MEMBER_SIZE:
                raise ValueError(_("A file in the archive is too large."))
            files.append((_zip_member_name(info), archive.read(info)))
    return sorted(files)


def read_directory(root):
    root = Path(root)
    return sorted(
        (path.relative_to(root).as_posix(), path.read_bytes())
        for path in root.rglob("*")
        if path.is_file() and _is_excel(path.relative_to(root).as_posix())
    )


def import_path(path):
    """Import a directory, a .zip archive or a single .xlsx file from disk."""
    path = Path(path)
    importer = Importer()
    if path.is_dir():
        return importer.import_named_files(read_directory(path))
    if path.suffix.lower() == ".zip":
        with path.open("rb") as fileobj:
            return importer.import_named_files(read_zip(fileobj))
    return importer.import_named_files([(path.name, path.read_bytes())])
