import io
import zipfile
from datetime import date

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import translation
from openpyxl import Workbook, load_workbook

from .importer import (
    Importer, context_from_path, parse_correspondence, parse_email_line, parse_person_line,
    read_zip,
)
from .models import Company, Contact, Country, Email, Industry

COMPANY_HEADER = [
    "Название компании", "Сайт", "Номера телефонов", "Почтовые адреса (e-mail)",
    "Сверка с анкетой", "Проверка сайта 05.10.2026",
    "E-mail с сайта (новые, которых не было в списке)", "Руководство и закупки (с сайта)",
    "E-mail из почты CTB (новые, которых не было в списке)",
    "Переписка в почте CTB по адресам из списка",
]


def xlsx(header, *rows):
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def sample_zip():
    auto_by = xlsx(
        COMPANY_HEADER,
        [
            "ACME MOTORS", "www.acme.by\nwww.acme-group.by", "375 17 111 11 11\n+375 29 222 22 22",
            "info@acme.by\nsales@acme.by", "Анкета 2024 г. Совпало: www.acme.by.",
            "Сайт прочитан.", "snab@acme.by — Отдел снабжения\nhr@acme.by",
            "Директор: Иванов Иван Иванович, Генеральный директор — e-mail не указан\n"
            "Закупки/снабжение: Петров Пётр, Начальник отдела снабжения — snab@acme.by\n"
            "Закупки/снабжение: Отдел закупок — zakupki@acme.by",
            "guess@acme.by — писем от нас: 2, посл. 13.05.2025; привязан по названию адреса, проверьте",
            "sales@acme.by (Сидоров С.) — писем от нас: 3, посл. 16.10.2024; "
            "от адреса: отвечал письмом (2), посл. 17.10.2024",
        ],
        [None, None, None, None, None, None, None, None, None, None],
        [None, "orphan.by", None, None, None, None, None, None, None, None],
    )
    appliances_ru = xlsx(COMPANY_HEADER, ["Beta Home", "beta.ru", None, "office@beta.ru"])
    unassigned = xlsx(
        ["Группа", "E-mail", "Имя (как в почте)", "Что есть в почте CTB", "Комментарий"],
        ["Другие страны", "info@other.kz", None, "писем от нас: 4, посл. 11.10.2024", "Казахстан"],
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("2026/Автомобильная/Беларусь/Автомобильная - Беларусь.xlsx", auto_by)
        archive.writestr("2026/Бытовая техника/РФ/Бытовая техника - РФ.xlsx", appliances_ru)
        archive.writestr("2026/Почта CTB - адреса вне списков.xlsx", unassigned)
        archive.writestr("2026/~$lock.xlsx", b"junk")
        archive.writestr("2026/readme.txt", b"ignored")
    buffer.seek(0)
    return buffer


class ParserTests(TestCase):
    def test_parse_email_line(self):
        self.assertEqual(
            parse_email_line("a.b@x.ru (Иванов И.) — писем от нас: 1"),
            ("a.b@x.ru", "Иванов И.", "писем от нас: 1"),
        )
        self.assertEqual(parse_email_line("Info@X.ru"), ("info@x.ru", "", ""))
        self.assertEqual(parse_email_line("no email here")[0], None)

    def test_parse_correspondence(self):
        parsed = parse_correspondence(
            "писем от нас: 4, посл. 13.05.2025; от адреса: отвечал письмом (2), посл. 13.05.2025")
        self.assertEqual(parsed["sent_count"], 4)
        self.assertEqual(parsed["last_sent_on"], date(2025, 5, 13))
        self.assertTrue(parsed["replied"])
        self.assertFalse(parsed["needs_review"])
        self.assertTrue(parse_correspondence("писем от нас: 1; проверьте")["needs_review"])

    def test_parse_person_line(self):
        self.assertEqual(
            parse_person_line("Директор: Иванов И.И., Генеральный директор ОАО «X» — e-mail не указан"),
            {"full_name": "Иванов И.И.", "position": "Генеральный директор ОАО «X»",
             "category": Contact.Category.MANAGEMENT, "email": ""},
        )
        parsed = parse_person_line("Закупки/снабжение: Отдел снабжения — snab@x.ru")
        self.assertEqual((parsed["full_name"], parsed["position"]), ("", "Отдел снабжения"))
        self.assertEqual(parsed["category"], Contact.Category.PROCUREMENT)
        self.assertEqual(parsed["email"], "snab@x.ru")

    def test_context_from_path(self):
        self.assertEqual(context_from_path("2026/Авто/РФ/list.xlsx"), ("Авто", "РФ"))
        self.assertEqual(context_from_path("Авто - РФ.xlsx"), ("Авто", "РФ"))
        self.assertEqual(context_from_path("list.xlsx"), (None, None))


class ImportTests(TestCase):
    def setUp(self):
        translation.activate("en")

    def run_import(self):
        return Importer().import_named_files(read_zip(sample_zip()))

    def test_import_structure(self):
        report = self.run_import()
        self.assertEqual(len(report.files), 3)
        self.assertEqual(report.companies_created, 2)
        self.assertEqual(report.unassigned_emails_created, 1)
        self.assertEqual(len(report.warnings), 1)  # the row without a company name

        acme = Company.objects.get(name="ACME MOTORS")
        self.assertEqual(acme.country.iso_code, "BY")
        self.assertEqual(acme.country.name_tr, "Belarus")
        self.assertEqual([str(i) for i in acme.industries.all()], ["Automotive"])
        self.assertEqual(acme.website_list, ["www.acme.by", "www.acme-group.by"])
        self.assertEqual(len(acme.phone_list), 2)
        self.assertIn("Анкета 2024", acme.questionnaire_note)

        emails = {e.email: e for e in acme.emails.all()}
        self.assertEqual(
            set(emails),
            {"info@acme.by", "sales@acme.by", "snab@acme.by", "hr@acme.by", "guess@acme.by",
             "zakupki@acme.by"},
        )
        self.assertEqual(emails["info@acme.by"].source, Email.Source.LIST)
        self.assertEqual(emails["snab@acme.by"].description, "Отдел снабжения")
        self.assertTrue(emails["sales@acme.by"].replied)
        self.assertEqual(emails["sales@acme.by"].sent_count, 3)
        self.assertEqual(emails["sales@acme.by"].person_name, "Сидоров С.")
        self.assertEqual(emails["guess@acme.by"].source, Email.Source.CTB_MAIL)
        self.assertTrue(emails["guess@acme.by"].needs_review)
        self.assertEqual(acme.contacts.count(), 3)

        beta = Company.objects.get(name="Beta Home")
        self.assertEqual(beta.country.iso_code, "RU")
        self.assertEqual(str(beta.industries.get()), "Home appliances")

        loose = Email.objects.get(company=None)
        self.assertEqual((loose.email, loose.group, loose.sent_count), ("info@other.kz", "Другие страны", 4))
        # The root-level file must not create a fake industry or country.
        self.assertEqual(Industry.objects.count(), 2)
        self.assertEqual(Country.objects.count(), 2)

    def test_reimport_is_idempotent(self):
        self.run_import()
        counts = (Company.objects.count(), Email.objects.count(), Contact.objects.count())
        report = self.run_import()
        self.assertEqual(report.companies_created, 0)
        self.assertEqual(report.companies_updated, 2)
        self.assertEqual(report.emails_created, 0)
        self.assertEqual(
            counts, (Company.objects.count(), Email.objects.count(), Contact.objects.count()))
        self.assertEqual(Company.objects.get(name="ACME MOTORS").website_list,
                         ["www.acme.by", "www.acme-group.by"])

    def test_same_name_in_another_country_is_another_company(self):
        self.run_import()
        Importer().import_named_files([("Автомобильная/Казахстан/x.xlsx",
                                        xlsx(COMPANY_HEADER, ["Acme  motors", "acme.kz"]))])
        self.assertEqual(Company.objects.filter(normalized_name="acme motors").count(), 2)

    def test_unknown_industry_is_created_with_warning(self):
        report = Importer().import_named_files(
            [("Космос/РФ/x.xlsx", xlsx(COMPANY_HEADER, ["Orbit", "orbit.ru"]))])
        self.assertEqual(Industry.objects.get().name_en, "Космос")
        self.assertEqual(len(report.warnings), 1)

    def test_broken_file_is_reported(self):
        report = Importer().import_named_files([("a/b/broken.xlsx", b"not excel")])
        self.assertEqual(Company.objects.count(), 0)
        self.assertEqual(len(report.warnings), 1)


class AdminTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_superuser("admin", "a@example.com", "pw")

    def setUp(self):
        translation.activate("en")
        self.client.force_login(self.user)

    def test_pages_render_in_both_languages(self):
        Importer().import_named_files(read_zip(sample_zip()))
        company = Company.objects.get(name="ACME MOTORS")
        urls = [
            reverse("admin:index"),
            reverse("admin:directory_company_changelist"),
            reverse("admin:directory_company_changelist") + "?replied=yes",
            reverse("admin:directory_company_change", args=[company.pk]),
            reverse("admin:directory_company_import_excel"),
            reverse("admin:directory_email_changelist"),
            reverse("admin:directory_contact_changelist"),
            reverse("admin:directory_country_changelist"),
            reverse("admin:directory_industry_changelist"),
            reverse("admin:directory_event_changelist"),
            reverse("admin:directory_meeting_changelist"),
        ]
        for language, marker in (("en", "Companies"), ("tr", "Firmalar")):
            self.client.cookies["django_language"] = language
            for url in urls:
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200, url)
            self.assertContains(self.client.get(reverse("admin:index")), marker)

    def test_country_name_follows_language(self):
        country = Country.objects.create(name_en="Russia", name_tr="Rusya", iso_code="ru")
        self.assertEqual(country.iso_code, "RU")
        with translation.override("tr"):
            self.assertEqual(str(country), "Rusya")
        with translation.override("en"):
            self.assertEqual(str(country), "Russia")

    def test_import_view_upload(self):
        upload = SimpleUploadedFile("2026.zip", sample_zip().getvalue())
        response = self.client.post(reverse("admin:directory_company_import_excel"), {"file": upload})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Company.objects.count(), 2)
        self.assertContains(response, "Companies: 2 new")

    def test_import_view_single_file_with_chosen_country(self):
        country = Country.objects.create(name_en="Türkiye", name_tr="Türkiye", iso_code="TR")
        upload = SimpleUploadedFile("list.xlsx", xlsx(COMPANY_HEADER, ["Gamma", "gamma.com.tr"]))
        self.client.post(reverse("admin:directory_company_import_excel"),
                         {"file": upload, "country": country.pk})
        self.assertEqual(Company.objects.get().country, country)

    def test_import_view_rejects_bad_files(self):
        response = self.client.post(reverse("admin:directory_company_import_excel"),
                                    {"file": SimpleUploadedFile("x.zip", b"not a zip")})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Company.objects.count(), 0)
        self.assertTrue(response.context["form"].errors)

    def test_export_xlsx(self):
        Importer().import_named_files(read_zip(sample_zip()))
        response = self.client.post(reverse("admin:directory_company_changelist"), {
            "action": "export_xlsx",
            "_selected_action": list(Company.objects.values_list("pk", flat=True)),
        })
        self.assertEqual(response.status_code, 200)
        ws = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws["A2"].value, "ACME MOTORS")
        self.assertIn("info@acme.by", ws["G2"].value)

    def test_exported_file_can_be_imported_back(self):
        Importer().import_named_files(read_zip(sample_zip()))
        response = self.client.post(reverse("admin:directory_company_changelist"), {
            "action": "export_xlsx",
            "_selected_action": list(Company.objects.values_list("pk", flat=True)),
        })
        Company.objects.all().delete()
        translation.activate("en")
        Importer().import_named_files([("export.xlsx", response.content)])
        acme = Company.objects.get(name="ACME MOTORS")
        self.assertEqual(acme.country.iso_code, "BY")
        self.assertEqual(str(acme.industries.get()), "Automotive")
        self.assertIn("info@acme.by", acme.emails.values_list("email", flat=True))


class CsrfFailureTests(TestCase):
    def test_reason_is_shown(self):
        client = self.client_class(enforce_csrf_checks=True)
        response = client.post(reverse("admin:login"), {"username": "x", "password": "y"})
        self.assertEqual(response.status_code, 403)
        self.assertContains(response, "CSRF cookie not set", status_code=403)

    def test_https_behind_proxy_passes_origin_check(self):
        get_user_model().objects.create_superuser("admin", "a@example.com", "pw")
        client = self.client_class(enforce_csrf_checks=True)
        headers = {"HTTP_HOST": "testserver", "HTTP_X_FORWARDED_PROTO": "https",
                   "HTTP_ORIGIN": "https://testserver",
                   "HTTP_REFERER": "https://testserver/admin/login/"}
        page = client.get(reverse("admin:login"), **headers)
        token = page.context["csrf_token"]
        response = client.post(reverse("admin:login"),
                               {"username": "admin", "password": "pw", "csrfmiddlewaretoken": token},
                               **headers)
        self.assertEqual(response.status_code, 302)
